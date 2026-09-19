"""Build an editable PowerPoint figure from the six unmodified PNG panels."""

from collections import Counter
import hashlib
import json
from pathlib import Path
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent
RENDERS = OUT / "renders"
OUTPUT = OUT / "reference_pipeline_editable.pptx"
# Same 10:3 canvas as the original figure. Native objects can also be copied
# together into the user's existing slide and resized as a group.
CANVAS_WIDTH = 4200
CANVAS_HEIGHT = 1260
SLIDE_WIDTH_INCHES = 20
EMU_PER_PIXEL = Inches(SLIDE_WIDTH_INCHES) / CANVAS_WIDTH
POINTS_PER_PIXEL = SLIDE_WIDTH_INCHES * 72 / CANVAS_WIDTH


def length(value):
    return round(value * EMU_PER_PIXEL)


def color(value):
    return RGBColor.from_string(value.lstrip("#"))


def main():
    prs = Presentation()
    prs.slide_width = length(CANVAS_WIDTH)
    prs.slide_height = length(CANVAS_HEIGHT)
    prs.core_properties.title = "Reference Motion Processing — Editable"
    prs.core_properties.subject = "Native editable labels, arrows and panels; original F7/B7 images"
    prs.core_properties.author = "YHL"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color("ffffff")

    def text(name, x, center_y, width, label, size, fill="162c49", bold=False, align=PP_ALIGN.CENTER):
        height = size * 1.65
        shape = slide.shapes.add_textbox(length(x), length(center_y - height / 2), length(width), length(height))
        shape.name = name
        frame = shape.text_frame
        frame.clear()
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.0
        run = paragraph.add_run()
        run.text = label
        run.font.name = "Arial"
        run.font.size = Pt(size * POINTS_PER_PIXEL)
        run.font.bold = bold
        run.font.color.rgb = color(fill)
        return shape

    def card(name, x):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, length(x), length(220), length(1100), length(896))
        shape.name = name
        shape.adjustments[0] = 24 / 896
        shape.fill.solid()
        shape.fill.fore_color.rgb = color("f6f9fd")
        shape.line.color.rgb = color("cfdded")
        shape.line.width = Pt(3 * POINTS_PER_PIXEL)

    def arrow(name, left, right, y, label, tool):
        center = (left + right) / 2
        text(name + " - Label", center - 170, y - 85, 340, label, 35, bold=True)
        shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, length(left), length(y - 21), length(right - left), length(42))
        shape.name = name + " - Editable arrow"
        shape.fill.solid()
        shape.fill.fore_color.rgb = color("3569ae")
        shape.line.fill.background()
        shape.adjustments[0] = 8 / 42
        shape.adjustments[1] = 33 / 42
        text(name + " - Method", center - 150, y + 64, 300, tool, 31, fill="5a7192")

    text("Figure title", 90, 88, 3000, "Reference Motion Processing", 64, bold=True, align=PP_ALIGN.LEFT)
    text("Figure subtitle", 90, 151, 3000, "Human demonstrations to robot reference motion", 32, fill="61728a", align=PP_ALIGN.LEFT)
    sources = []
    for x, title, suffix in [(90, "Human Video", "video"), (1550, "Reconstructed Human", "reconstruction"), (3010, "Retargeted Robot", "robot")]:
        card(title + " - Editable panel", x)
        text(title + " - Heading", x + 30, 263, 1040, title, 39, bold=True)
        for i, (code, stroke) in enumerate([("F7", "Forehand"), ("B7", "Backhand")]):
            left = x + 30 + i * 540
            path = RENDERS / f"{code}_{suffix}.png"
            picture = slide.shapes.add_picture(str(path), length(left), length(310), width=length(510), height=length(756))
            picture.name = f"{title} - {stroke} ({code}) - Original PNG"
            text(f"{title} - {stroke} caption", left, 1090, 510, stroke, 29, fill="476381")
            sources.append(path)
    arrow("Reconstruction", 1220, 1520, 665, "Reconstruction", "GVHMR")
    arrow("Retargeting", 2680, 2980, 665, "Retargeting", "GMR")
    text("NPZ post-processing note", 3010, 1174, 1100, "Post-processing & alignment → Reference NPZ", 30, fill="476381")

    slide.notes_slide.notes_text_frame.text = (
        "Editable version for YHL. Six separate PNG pictures are embedded without changing their bytes. "
        "All 16 labels are native text boxes; the 3 panels and 2 arrows are native PowerPoint shapes. "
        "Use the Selection Pane to select objects by their descriptive names. "
        "To insert the complete diagram into another presentation: select all slide objects, copy and paste, "
        "then group and resize proportionally if needed. Ungroup to edit individually. "
        "This slide retains the original 10:3 figure aspect ratio. "
        "Dataset: 20260806_mqy_capture. F7: forehand clip 010, video/SMPL-X/pre-align frame 34, aligned frame 31. "
        "B7: backhand clip 011, video/SMPL-X/pre-align frame 35, aligned frame 31. "
        "Robot panels are PyBullet renders of the matching NPZ poses."
    )
    prs.save(OUTPUT)

    # Reopen the saved package and validate actual object types, not just generation inputs.
    reopened = Presentation(OUTPUT)
    assert len(reopened.slides) == 1
    objects = list(reopened.slides[0].shapes)
    pictures = [s for s in objects if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    texts = [s for s in objects if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    vectors = [s for s in objects if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert len(pictures) == 6 and len(texts) == 16 and len(vectors) == 5
    expected_hashes = Counter(hashlib.sha256(p.read_bytes()).hexdigest() for p in sources)
    actual_hashes = Counter(hashlib.sha256(p.image.blob).hexdigest() for p in pictures)
    assert expected_hashes == actual_hashes, "Embedded images changed"
    assert sum(s.auto_shape_type == MSO_SHAPE.RIGHT_ARROW for s in vectors) == 2
    assert sum(s.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE for s in vectors) == 3
    for shape in objects:
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= reopened.slide_width
        assert shape.top + shape.height <= reopened.slide_height
    with zipfile.ZipFile(OUTPUT) as package:
        assert package.testzip() is None
    report = {"file": str(OUTPUT), "slide_count": 1, "picture_count": 6, "editable_text_boxes": 16,
              "editable_arrows": 2, "editable_panels": 3, "all_images_byte_identical": True,
              "canvas_aspect_ratio": "10:3", "font": "Arial", "object_count": len(objects)}
    (OUT / "editable_pptx_validation.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
